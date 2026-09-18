#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
按征稿格式，从 Markdown 正文生成可下载的 Word。

格式（固化在 word_format.py）：
  A4；题目三号黑体居中；摘要/关键词/正文/参考文献小四宋体、1.5 倍行距；
  顺序为目录→题目→摘要→关键词→正文→参考文献；文中不署单位和姓名。

用法：
  python3 scripts/build_formatted_docx.py
  python3 scripts/build_formatted_docx.py --md 论文.md --out output/论文.docx
"""

from __future__ import annotations

import argparse
import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from word_format import SPEC  # noqa: E402

# 正文层级仅允许 一、／（一）／1.；不再把“缘起。”等段首词当标题加粗。
LEAD_WORDS = ()


def set_run_font(run, cn: str, size: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = SPEC.font_en
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), SPEC.font_en)
    rFonts.set(qn("w:hAnsi"), SPEC.font_en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), SPEC.font_en)


def apply_line_15(pf) -> None:
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.line_spacing = SPEC.line_spacing


def set_style(style, cn, size, bold, align, first_indent, space_before=0, space_after=0, outline=None):
    style.font.name = SPEC.font_en
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), SPEC.font_en)
    rFonts.set(qn("w:hAnsi"), SPEC.font_en)
    rFonts.set(qn("w:eastAsia"), cn)
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    apply_line_15(pf)
    pf.first_line_indent = first_indent
    if outline is not None:
        pPr = style.element.get_or_add_pPr()
        lvl = pPr.find(qn("w:outlineLvl"))
        if lvl is None:
            lvl = OxmlElement("w:outlineLvl")
            pPr.append(lvl)
        lvl.set(qn("w:val"), str(outline))


def init_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(SPEC.page_width_cm)
    sec.page_height = Cm(SPEC.page_height_cm)
    sec.left_margin = Cm(SPEC.margin_left_cm)
    sec.right_margin = Cm(SPEC.margin_right_cm)
    sec.top_margin = Cm(SPEC.margin_top_cm)
    sec.bottom_margin = Cm(SPEC.margin_bottom_cm)
    sec.header.paragraphs[0].text = ""
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)
    apply_line_15(fp.paragraph_format)
    run = fp.add_run()
    set_run_font(run, SPEC.font_cn_song, SPEC.body_pt)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)

    set_style(doc.styles["Normal"], SPEC.font_cn_song, SPEC.body_pt, False,
              WD_ALIGN_PARAGRAPH.JUSTIFY, Pt(SPEC.first_line_indent_pt))
    set_style(doc.styles["Title"], SPEC.font_cn_hei, SPEC.title_pt, True,
              WD_ALIGN_PARAGRAPH.CENTER, Cm(0), space_before=12, space_after=12)
    set_style(doc.styles["Heading 1"], SPEC.font_cn_hei, SPEC.h1_pt, True,
              WD_ALIGN_PARAGRAPH.CENTER, Cm(0), space_before=12, space_after=6, outline=0)
    set_style(doc.styles["Heading 2"], SPEC.font_cn_hei, SPEC.h2_pt, True,
              WD_ALIGN_PARAGRAPH.LEFT, Cm(0), space_before=8, space_after=4, outline=1)
    set_style(doc.styles["Heading 3"], SPEC.font_cn_hei, SPEC.body_pt, True,
              WD_ALIGN_PARAGRAPH.LEFT, Cm(0), space_before=6, space_after=3, outline=2)
    return doc


def add_para(doc, text, *, style="Normal", cn=None, size=None, bold=False,
             align=None, first_indent=None, space_before=0, space_after=0,
             keep_with_next=False):
    cn = cn or SPEC.font_cn_song
    size = SPEC.body_pt if size is None else size
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if first_indent is not None:
        pf.first_line_indent = first_indent
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    apply_line_15(pf)
    pf.widow_control = True
    pf.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, cn, size, bold)
    return p


def add_mixed(doc, parts, *, first_indent=None, space_before=0, space_after=0, align=None):
    """parts: list of (text, cn, bold)."""
    p = doc.add_paragraph(style="Normal")
    p.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(SPEC.first_line_indent_pt) if first_indent is None else first_indent
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    apply_line_15(pf)
    for text, cn, bold in parts:
        if text:
            run = p.add_run(text)
            set_run_font(run, cn, SPEC.body_pt, bold)
    return p


def set_cell_border(cell, **edges):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, (val, sz) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        old = tcBorders.find(qn(f"w:{edge}"))
        if old is not None:
            tcBorders.remove(old)
        tcBorders.append(el)


def add_three_line_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        tblBorders.append(el)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tblBorders)

    n_rows = len(rows)
    for ri, row in enumerate(rows):
        if ri == 0:
            tr = table.rows[0]._tr
            trPr = tr.get_or_add_trPr()
            if trPr.find(qn("w:tblHeader")) is None:
                hdr = OxmlElement("w:tblHeader")
                trPr.append(hdr)
        for ci in range(ncols):
            text = row[ci].strip() if ci < len(row) else ""
            cell = table.cell(ri, ci)
            cell.text = text
            is_header = ri == 0
            edges = {"left": ("nil", 0), "right": ("nil", 0), "top": ("nil", 0), "bottom": ("nil", 0)}
            if ri == 0:
                edges["top"] = ("single", 12)
                edges["bottom"] = ("single", 6)
            elif ri == n_rows - 1:
                edges["bottom"] = ("single", 12)
            set_cell_border(cell, **edges)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.first_line_indent = Cm(0)
                apply_line_15(para.paragraph_format)
                for run in para.runs:
                    set_run_font(run, SPEC.font_cn_hei if is_header else SPEC.font_cn_song,
                                 SPEC.table_pt, bold=is_header)


def parse_img_opts(raw: str | None) -> dict:
    opts = {"w": SPEC.image_width_cm, "hmax": 18.0}
    if not raw:
        return opts
    for part in raw.split(","):
        if "=" not in part:
            continue
        k, v = part.split("=", 1)
        k, v = k.strip(), v.strip()
        if k in ("w", "hmax"):
            try:
                opts[k] = float(v)
            except ValueError:
                pass
    return opts


def add_image(doc, path: Path, *, width_cm: float | None = None, hmax_cm: float = 18.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    apply_line_15(pf)
    run = p.add_run()
    width_cm = SPEC.image_width_cm if width_cm is None else width_cm
    try:
        from PIL import Image as PILImage
        with PILImage.open(path) as im:
            w, h = im.size
        if w and h:
            height_cm = width_cm * h / w
            if height_cm > hmax_cm:
                width_cm = hmax_cm * w / h
    except Exception:
        pass
    run.add_picture(str(path), width=Cm(width_cm))


def strip_md(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^#{1,6}\s*", "", text)
    text = re.sub(r"^\*\*(.+)\*\*$", r"\1", text)
    text = text.replace("**", "")
    return text.strip()


def is_blocked(text: str) -> bool:
    if not SPEC.omit_author:
        return False
    compact = re.sub(r"\s+", "", text)
    if compact in {"目录", "三级目录"}:
        return True
    for token in SPEC.blocked_signatures:
        if token in compact and len(compact) <= 20:
            return True
    return False


def split_table_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.replace("<br>", "\n").strip() for c in line.split("|")]


def is_sep_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells if c)


def collect_headings(lines: list[str]) -> list[tuple[int, str]]:
    items: list[tuple[int, str]] = []
    for line in lines:
        s = line.strip()
        if s.startswith("#### "):
            lv, raw = 3, s[5:]
        elif s.startswith("### "):
            lv, raw = 2, s[4:]
        elif s.startswith("## "):
            lv, raw = 1, s[3:]
        else:
            continue
        title = strip_md(raw)
        if title in {"目录", "三级目录"}:
            continue
        items.append((lv, title))
    return items


def add_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    add_para(doc, "目  录", cn=SPEC.font_cn_hei, size=SPEC.h1_pt, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0),
             space_before=0, space_after=6)
    for lv, title in headings:
        cn = SPEC.font_cn_hei if lv == 1 else SPEC.font_cn_song
        p = add_para(doc, title, cn=cn, size=SPEC.body_pt, bold=(lv == 1),
                     align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0),
                     space_before=0, space_after=0)
        p.paragraph_format.left_indent = Cm(0.74 * (lv - 1))
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
    br = doc.add_paragraph()
    br.paragraph_format.first_line_indent = Cm(0)
    br.paragraph_format.space_before = Pt(0)
    br.paragraph_format.space_after = Pt(0)
    br.add_run().add_break(WD_BREAK.PAGE)


def parse_and_build(md_path: Path, out_path: Path) -> Path:
    raw = md_path.read_text(encoding="utf-8")
    lines = raw.splitlines()
    base = md_path.parent
    doc = init_document()
    if doc.paragraphs and not doc.paragraphs[0].text.strip():
        el = doc.paragraphs[0]._element
        el.getparent().remove(el)
    add_toc(doc, collect_headings(lines))

    i = 0
    skipping_toc = False
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped in ("## 三级目录", "## 目录", "三级目录"):
            skipping_toc = True
            i += 1
            continue
        if skipping_toc:
            if stripped.startswith("#") or stripped.startswith("【"):
                skipping_toc = False
            else:
                i += 1
                continue

        text = strip_md(stripped)
        if not text or is_blocked(text):
            i += 1
            continue

        # 图片，可选 {w=14.5,hmax=18}
        m_img = re.match(r"!\[.*?\]\((.+?)\)(?:\{([^}]+)\})?", stripped)
        if m_img:
            img = (base / m_img.group(1)).resolve()
            if img.exists():
                opts = parse_img_opts(m_img.group(2))
                add_image(doc, img, width_cm=opts["w"], hmax_cm=opts["hmax"])
            else:
                print("缺图，已跳过：", img)
            i += 1
            continue

        # Markdown 表格
        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = split_table_row(lines[i])
                if not is_sep_row(cells):
                    rows.append(cells)
                i += 1
            add_three_line_table(doc, rows)
            continue

        # 题目 / 副标题
        if stripped.startswith("# ") and not stripped.startswith("## "):
            is_sub = text.startswith("——") or text.startswith("—")
            add_para(doc, text, style="Title", cn=SPEC.font_cn_hei, size=SPEC.title_pt,
                     bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0),
                     space_before=0 if is_sub else 12, space_after=12 if is_sub else 0)
            i += 1
            continue

        # 摘要 / 关键词
        if text.startswith("【摘要】"):
            add_mixed(doc, [("【摘要】", SPEC.font_cn_hei, True),
                            (text[len("【摘要】"):], SPEC.font_cn_song, False)],
                      space_before=6, space_after=6)
            i += 1
            continue
        if text.startswith("【关键词】"):
            add_mixed(doc, [("【关键词】", SPEC.font_cn_hei, True),
                            (text[len("【关键词】"):], SPEC.font_cn_song, False)],
                      space_before=0, space_after=12)
            i += 1
            continue

        # 一级 / 参考文献
        if stripped.startswith("## "):
            add_para(doc, text, style="Heading 1", cn=SPEC.font_cn_hei, size=SPEC.h1_pt, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0),
                     space_before=12, space_after=6)
            i += 1
            continue
        if stripped.startswith("### "):
            add_para(doc, text, style="Heading 2", cn=SPEC.font_cn_hei, size=SPEC.h2_pt, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0),
                     space_before=8, space_after=4)
            i += 1
            continue
        if stripped.startswith("#### "):
            add_para(doc, text, style="Heading 3", cn=SPEC.font_cn_hei, size=SPEC.body_pt, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0),
                     space_before=6, space_after=3)
            i += 1
            continue

        # 图题 / 表题（统一“图X 标题 / 表X 标题”）
        if re.match(r"^[图表]\d+", text):
            text = re.sub(r"^([图表]\d+)[　\s]+", r"\1 ", text)
            add_para(doc, text, cn=SPEC.font_cn_song, size=SPEC.caption_pt,
                     align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0),
                     space_before=2, space_after=8)
            i += 1
            continue

        # 案例举隅
        if text.startswith("【案例举隅"):
            add_para(doc, text, cn=SPEC.font_cn_hei, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0),
                     space_before=8, space_after=4, keep_with_next=True)
            i += 1
            continue

        # 参考文献条目
        if re.match(r"^\[\d+\]", text):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent = Cm(0.74)
            pf.first_line_indent = Cm(-0.74)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            apply_line_15(pf)
            run = p.add_run(text)
            set_run_font(run, SPEC.font_cn_song, SPEC.ref_pt)
            i += 1
            continue

        lead = next((w for w in LEAD_WORDS if text.startswith(w)), None)
        if lead:
            add_mixed(doc, [(lead, SPEC.font_cn_hei, True),
                            (text[len(lead):], SPEC.font_cn_song, False)])
        else:
            add_para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                     first_indent=Pt(SPEC.first_line_indent_pt))
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def default_md() -> Path:
    cands = sorted(ROOT.glob("*三元协同*.md")) + sorted(ROOT.glob("*.md"))
    if not cands:
        raise SystemExit("未找到 Markdown 正文")
    return cands[0]


def main() -> None:
    parser = argparse.ArgumentParser(description="按征稿格式生成 Word")
    parser.add_argument("--md", type=Path, default=None, help="Markdown 正文路径")
    parser.add_argument("--out", type=Path, default=None, help="输出 docx 路径")
    args = parser.parse_args()
    md = args.md.resolve() if args.md else default_md()
    out = args.out
    if out is None:
        out = ROOT / "output" / f"{md.stem}.docx"
    else:
        out = out.resolve()
    path = parse_and_build(md, out)
    ascii_out = path.parent / "paper.docx"
    ascii_out.write_bytes(path.read_bytes())
    root_ascii = ROOT / "paper.docx"
    root_ascii.write_bytes(path.read_bytes())
    artifacts = Path("/opt/cursor/artifacts")
    if artifacts.exists():
        (artifacts / "paper.docx").write_bytes(path.read_bytes())
        print("artifacts:", artifacts / "paper.docx")
    print("已生成符合征稿格式的 Word：")
    print(path)
    print("英文文件名（请下载这个）：")
    print(ascii_out)
    print(root_ascii)


if __name__ == "__main__":
    main()
